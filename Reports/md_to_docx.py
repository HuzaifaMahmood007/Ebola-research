"""md_to_docx.py -- render the Week-4 results markdown reports as .docx.

WHY A CONVERTER AND NOT A DOCUMENT BUILDER. The numbers in the markdown reports have already been
verified cell-by-cell against the run JSONs (264 assertions, doc -> data). Hand-building a parallel
.docx from the same data would re-open every transcription risk that verification just closed, and
the two artifacts would drift the moment one is regenerated. So the markdown is the single source of
truth and this only re-renders it. If a number is wrong here, it is wrong in the markdown, and the
markdown is what gets checked.

Deliberately a SMALL subset of markdown -- exactly what these two reports use: ATX headings, pipe
tables, `-` bullets, blockquotes, fenced code, horizontal rules, and inline `**bold**`/`code`.
Anything else passes through as plain text rather than failing, because a converter that crashes on
an unexpected character is worse than one that renders it literally.

Run with BASE python (that is where python-docx lives), from anywhere:
    python Reports/md_to_docx.py
    python Reports/md_to_docx.py somefile.md      # single file
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches

HERE = Path(__file__).resolve().parent
DEFAULT = ["Encoder_Results_Consolidated.md", "Encoder_Results_Stakeholder_Brief.md"]

# House style, matched to Reports/Week3_Results_Summary.docx (see make_results_doc.py):
# Calibri 11 body, "Light Grid Accent 1" tables, bold 9pt header row, 9pt body cells,
# level-0 title, no manual colours or cell shading.
BODY_PT = Pt(11)
CELL_PT = Pt(9)
TABLE_STYLE = "Light Grid Accent 1"

# inline: **bold**, *italic*, `code`, [text](target). Applied to a JOINED paragraph, never to a
# single source line -- markdown wraps mid-span, so line-at-a-time parsing leaves stray asterisks.
_INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\s)[^*]+?\*|`[^`]+`|\[[^\]]+\]\([^)]*\))")


def add_runs(par, text, bold=False, italic=False):
    """Render one logical paragraph into bold / italic / code / plain runs.

    RECURSIVE, because spans nest: `**`nrmse` is missing**` is a bold span containing code.
    Handling only the outer span would emit the literal backticks inside the bold run, which is
    exactly the leftover-markdown symptom this converter exists to avoid.
    """
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_runs(par, piece[2:-2], bold=True, italic=italic)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.bold, r.italic = bold, italic
        elif piece.startswith("[") and "](" in piece:
            r = par.add_run(piece[1:piece.index("]")])
            r.bold, r.italic = bold, True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            add_runs(par, piece[1:-1], bold=bold, italic=True)
        else:
            r = par.add_run(piece)
            r.bold, r.italic = bold, italic


def repeat_header(t):
    """Repeat the header row when a table breaks across pages."""
    pr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def is_table_row(ln):
    return ln.startswith("|") and ln.endswith("|")


def is_divider(ln):
    return bool(re.match(r"^\|[\s\-:|]+\|$", ln))


def emit_table(doc, rows):
    """rows[0] is the header. Styling matches Week3_Results_Summary.docx exactly."""
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    try:
        t.style = TABLE_STYLE
    except KeyError:                       # style absent from a bare template; borders are cosmetic
        pass
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            par = cells[j].paragraphs[0]
            add_runs(par, row[j] if j < len(row) else "")
            for run in par.runs:
                run.font.size = CELL_PT
                if i == 0:
                    run.font.bold = True
    repeat_header(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path, out_path):
    """Render markdown to .docx.

    THE CENTRAL RULE: markdown hard-wraps a paragraph across several source lines, so every block is
    accumulated and only rendered once the block ENDS (blank line, heading, table, list, rule). Doing
    inline parsing per source line -- the first version of this file -- both splits sentences into
    separate Word paragraphs and leaves stray `**` wherever a bold span crossed a line break.
    """
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = BODY_PT

    buf, in_code, code, para, first_heading = [], False, [], [], True

    def flush_table():
        if buf:
            emit_table(doc, list(buf))
            buf.clear()

    def flush_para():
        """Join the wrapped source lines into ONE logical paragraph, then parse inline spans."""
        if not para:
            return
        text = " ".join(x.strip() for x in para).strip()
        para.clear()
        if not text:
            return
        m = re.match(r"^[-*]\s+(.*)$", text)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            return
        m = re.match(r"^(\d+)\.\s+(.*)$", text)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(2))
            return
        if text.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, text[2:])
            for r in p.runs:
                r.italic = True
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, text)

    for raw in lines:
        ln = raw.rstrip()
        s = ln.strip()

        if s.startswith("```"):
            flush_para()
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                r = p.add_run("\n".join(code))
                r.font.name, r.font.size = "Consolas", Pt(8.5)
                code.clear()
            in_code = not in_code
            continue
        if in_code:
            code.append(ln)
            continue

        if is_table_row(s):
            flush_para()
            if not is_divider(s):
                buf.append([c.strip() for c in s.strip("|").split("|")])
            continue
        flush_table()

        if not s:                                   # blank line ends the logical paragraph
            flush_para()
            continue
        if re.match(r"^-{3,}$|^\*{3,}$|^_{3,}$", s):
            flush_para()                             # house style has no rules; headings carry structure
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            flush_para()
            title = re.sub(r"\*\*|`", "", m.group(2))
            lvl = len(m.group(1))
            if lvl == 1 and first_heading:           # document title, not a section heading
                doc.add_heading(title, level=0)
                first_heading = False
            else:
                doc.add_heading(title, level=min(lvl, 4))
            continue
        # a bullet or numbered item starts its own block, but may itself wrap over later lines
        if re.match(r"^([-*]|\d+\.)\s+", s) and para:
            flush_para()
        para.append(s)

    flush_para()
    flush_table()
    doc.save(out_path)
    return out_path


def _demo():
    """Runnable checks. The load-bearing one is WRAPPED MARKUP: the first version of this converter
    parsed line-at-a-time, so a bold span crossing a source line break leaked literal `**` into the
    document and split one sentence into two Word paragraphs. Both are asserted against here."""
    import tempfile, os
    md = ("# Doc Title\n\n"
          "## A section\n\n"
          "Across all 16 cells: **12 are significantly worse, 4 show no\n"
          "measurable effect, and none is better.** That is the finding.\n\n"
          "Some *italic* and `code` and a [link](x.md).\n\n"
          "**`nrmse` is missing here** — bold wrapping a code span.\n\n"
          "| a | b |\n|---|---|\n| 1.5 | -2.5 |\n\n"
          "- bullet that also wraps\n  onto a second line\n\n"
          "---\n\n1. numbered item\n")
    d = Path(tempfile.mkdtemp())
    src = d / "t.md"
    src.write_text(md, encoding="utf-8")
    out = convert(src, d / "t.docx")
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    body = "\n".join(texts)

    # 1. NO literal markdown may reach the document
    for junk in ("**", "`", "](", "###"):
        assert junk not in body, f"literal {junk!r} leaked into the docx: {body!r}"
    assert not any(re.match(r"^_{5,}$", t.strip()) for t in texts), "underscore rules must be gone"

    # 2. a bold span crossing a line break must become ONE paragraph with a real bold run
    joined = [t for t in texts if t.startswith("Across all 16 cells")]
    assert len(joined) == 1, f"wrapped paragraph must not be split: {joined}"
    assert "none is better." in joined[0] and joined[0].endswith("That is the finding."), joined
    bolds = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert any("12 are significantly worse" in b and "none is better" in b for b in bolds), \
        f"the wrapped bold span must be one bold run, got {bolds}"

    # 3. structure survives
    assert texts[0] == "Doc Title" and doc.paragraphs[0].style.name in ("Title", "Heading 1")
    assert len(doc.tables) == 1, "the pipe table must become a real table, not text"
    tb = doc.tables[0]
    assert tb.rows[0].cells[0].text == "a" and tb.rows[1].cells[1].text == "-2.5", \
        "table cells must survive with their signs intact"
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs), "bullets must survive"
    assert any(p.style.name == "List Number" for p in doc.paragraphs), "numbered items must survive"
    assert any("bullet that also wraps onto a second line" in t for t in texts), \
        "a wrapped bullet must rejoin into one item"
    assert any(r.italic for p in doc.paragraphs for r in p.runs), "italic must survive"
    # nested span: the code run inside a bold span must be BOTH bold and monospaced, with no backticks
    nested = [r for p in doc.paragraphs for r in p.runs if r.text == "nrmse"]
    assert nested and nested[0].bold and nested[0].font.name == "Consolas", \
        "a code span inside bold must stay bold AND monospaced, not emit literal backticks"
    assert not any("|" in t for t in texts), "no raw pipe rows may leak as text"
    os.remove(out)
    print("ok  no literal markdown leaks; wrapped bold/bullets rejoin into single paragraphs; "
          "title, tables with signed cells, lists and italics all survive")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if "--selfcheck" in sys.argv:
        _demo()
    else:
        for name in (args or DEFAULT):
            p = Path(name)
            if not p.is_absolute() and not p.exists():
                p = HERE / name
            out = p.with_suffix(".docx")
            convert(p, out)
            print(f"ok  {p.name} -> {out.name}  ({out.stat().st_size:,} bytes)")
